"""Clinical phenotype extraction from the challenge DOCX (text + tables)."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from rareai.log import get_logger

logger = get_logger(__name__)


def extract_docx_text(docx_path: Path | str) -> str:
    """Extract paragraphs and tables from the clinical phenotype document."""
    doc = Document(str(docx_path))
    lines: list[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table_idx, table in enumerate(doc.tables, start=1):
        lines.append(f"\n[TABLE {table_idx}]")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    return "\n".join(lines)


def save_phenotype_text(docx_path: Path | str, out_path: Path | str) -> Path:
    """Persist extracted text for curation (stays inside private data/)."""
    text = extract_docx_text(docx_path)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8")
    logger.info("Phenotype text extracted: %s (%d chars)", out, len(text))
    return out
